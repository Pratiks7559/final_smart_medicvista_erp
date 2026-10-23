from datetime import date
from io import BytesIO
import json
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from django.test import RequestFactory, SimpleTestCase, TestCase

from core import inventory_list2_views
from core.fast_inventory import FastInventory
from core.inventory_date_filters import get_inventory_date_range
from core.models import (
    CustomerMaster,
    InvoiceMaster,
    ProductMaster,
    PurchaseMaster,
    SalesInvoiceMaster,
    SalesMaster,
    SupplierMaster,
)
from core.views import inventory_search_suggestions


class InventoryDateRangeTests(SimpleTestCase):
    def test_parses_optional_inclusive_date_bounds_and_legacy_parameters(self):
        request = RequestFactory().get(
            '/reports/dateexpiry/',
            {'expiry_from': '2026-06-01', 'expiry_to': '2026-06-30'},
        )

        start, end, is_valid = get_inventory_date_range(
            request,
            legacy_start=('expiry_from',),
            legacy_end=('expiry_to',),
        )

        self.assertEqual((start, end, is_valid), (date(2026, 6, 1), date(2026, 6, 30), True))

    def test_rejects_invalid_or_reversed_dates(self):
        invalid_request = RequestFactory().get('/reports/batches/', {'date_from': 'not-a-date'})
        reversed_request = RequestFactory().get(
            '/reports/batches/',
            {'date_from': '2026-07-01', 'date_to': '2026-06-30'},
        )

        with patch('core.inventory_date_filters.messages.error'):
            self.assertEqual(get_inventory_date_range(invalid_request), (None, None, False))
            self.assertEqual(
                get_inventory_date_range(reversed_request),
                (date(2026, 7, 1), date(2026, 6, 30), False),
            )


class ProductAutocompleteTests(TestCase):
    def test_suggestions_match_partial_product_name_and_include_display_details(self):
        ProductMaster.objects.create(
            product_name='Test Medicine',
            product_company='Test Pharma',
            product_packing='10 tabs',
            product_salt='Test Salt',
            product_category='Tablet',
            product_hsn='3004',
            product_hsn_percent='5',
        )
        request = RequestFactory().get('/api/inventory-search-suggestions/', {'q': 'medic'})
        request.user = MagicMock(is_authenticated=True)

        response = inventory_search_suggestions(request)

        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertEqual(payload['products'][0]['product_name'], 'Test Medicine')
        self.assertEqual(payload['products'][0]['product_company'], 'Test Pharma')
        self.assertEqual(payload['suggestions'], ['Test Medicine - Test Pharma'])


class InventoryPeriodMovementTests(TestCase):
    def test_period_sales_show_negative_movement_from_earlier_purchase(self):
        supplier = SupplierMaster.objects.create(
            supplier_name='Test Supplier',
            supplier_mobile='9000000000',
        )
        product = ProductMaster.objects.create(
            product_name='Test Medicine',
            product_company='Test Pharma',
            product_packing='10 tabs',
            product_salt='Test Salt',
            product_category='Tablet',
            product_hsn='3004',
            product_hsn_percent='5',
        )
        old_invoice = InvoiceMaster.objects.create(
            invoice_no='OLD-001',
            invoice_date=date(2025, 1, 10),
            supplierid=supplier,
            transport_charges=0,
            invoice_total=100,
        )
        PurchaseMaster.objects.create(
            product_supplierid=supplier,
            product_invoiceid=old_invoice,
            product_invoice_no='OLD-001',
            productid=product,
            product_name=product.product_name,
            product_company=product.product_company,
            product_packing=product.product_packing,
            product_batch_no='BATCH-1',
            product_expiry='01-2030',
            product_MRP=12,
            product_purchase_rate=8,
            product_quantity=10,
            product_discount_got=0,
            product_transportation_charges=0,
        )
        customer = CustomerMaster.objects.create(customer_name='Test Customer')
        sale_invoice = SalesInvoiceMaster.objects.create(
            sales_invoice_no='SALE-001',
            sales_invoice_date=date(2026, 6, 15),
            customerid=customer,
        )
        SalesMaster.objects.create(
            sales_invoice_no=sale_invoice,
            customerid=customer,
            productid=product,
            product_name=product.product_name,
            product_company=product.product_company,
            product_packing=product.product_packing,
            product_batch_no='BATCH-1',
            product_expiry='01-2030',
            sale_quantity=13,
        )

        start_date, end_date = date(2026, 6, 1), date(2026, 6, 15)
        batches = FastInventory.get_batch_inventory_data(
            start_date=start_date,
            end_date=end_date,
        )
        expiry_groups, _ = FastInventory.get_dateexpiry_inventory_data(
            start_date=start_date,
            end_date=end_date,
        )

        self.assertEqual(len(batches), 1)
        self.assertEqual(batches[0]['stock'], -13)
        self.assertEqual(batches[0]['batch_no'], 'BATCH-1')
        expiry_batch = expiry_groups[0]['products'][0]
        self.assertEqual(expiry_batch['quantity'], -13)
        self.assertEqual(expiry_batch['batch_no'], 'BATCH-1')

        inventory_rows = inventory_list2_views._build_inventory2_data(
            start_date=start_date,
            end_date=end_date,
        )
        self.assertEqual(len(inventory_rows), 1)
        self.assertEqual(inventory_rows[0]['total_stock'], -13)

        inventory_request = RequestFactory().get(
            '/inventory2/',
            {
                'start_date': '2026-06-01',
                'end_date': '2026-06-15',
            },
        )
        inventory_request.user = MagicMock(is_authenticated=True)
        inventory_request.session = {'selected_year': 1990}
        with patch.object(inventory_list2_views, 'render', return_value=object()) as render:
            inventory_list2_views.inventory_list2.__wrapped__(inventory_request)
        inventory_context = render.call_args.args[2]
        self.assertEqual(inventory_context['inventory_data'][0]['total_stock'], -13)
        self.assertTrue(inventory_context['period_filtered'])

        history_request = RequestFactory().get(
            '/inventory2/history/',
            {'start_date': '2026-06-01', 'end_date': '2026-06-15'},
        )
        history_request.user = MagicMock(is_authenticated=True)
        with patch.object(inventory_list2_views, 'render', return_value=object()) as render:
            inventory_list2_views.inventory_transaction_history.__wrapped__(
                history_request, product.productid
            )
        history_rows = render.call_args.args[2]['transactions']
        self.assertEqual(len(history_rows), 1)
        self.assertEqual(history_rows[0]['transaction_display_date'], date(2026, 6, 15))
        self.assertEqual(history_rows[0]['balance'], -13)

        excel_request = RequestFactory().get(
            '/export/inventory2/excel/',
            {'start_date': '2026-06-01', 'end_date': '2026-06-15'},
        )
        excel_response = inventory_list2_views.export_inventory2_excel.__wrapped__(excel_request)
        self.assertEqual(excel_response.status_code, 200)
        worksheet = load_workbook(BytesIO(excel_response.content), data_only=True).active
        product_row = next(
            row for row in worksheet.iter_rows(values_only=True)
            if row[0] == product.product_name
        )
        self.assertEqual(product_row[5], -13)

        pdf_request = RequestFactory().get(
            '/export/inventory2/pdf/',
            {'start_date': '2026-06-01', 'end_date': '2026-06-15'},
        )
        pdf_response = inventory_list2_views.export_inventory2_pdf.__wrapped__(pdf_request)
        self.assertEqual(pdf_response.status_code, 200)
